from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from sqlalchemy import func
from datetime import datetime, date
from typing import Optional
from database import get_db
from models import Work, User
from auth import require_manager

router = APIRouter(prefix="/export", tags=["export"])


@router.get("/excel")
def export_excel(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    worker_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_manager),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from fastapi.responses import StreamingResponse
    import io

    q = db.query(Work).join(User, Work.worker_id == User.id)
    if date_from:
        q = q.filter(Work.start_time >= datetime.combine(date_from, datetime.min.time()))
    if date_to:
        q = q.filter(Work.start_time <= datetime.combine(date_to, datetime.max.time()))
    if worker_id:
        q = q.filter(Work.worker_id == worker_id)

    works = q.order_by(Work.start_time.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Работы"

    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = ["№", "Рабочий", "Номер детали", "Количество", "Начало", "Окончание", "Время (мин)", "Примечание"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for i, w in enumerate(works, 1):
        duration = ""
        if w.end_time and w.start_time:
            duration = round((w.end_time - w.start_time).total_seconds() / 60, 1)

        row_data = [
            i,
            w.worker.full_name,
            w.part_number,
            w.quantity,
            w.start_time.strftime("%d.%m.%Y %H:%M"),
            w.end_time.strftime("%d.%m.%Y %H:%M") if w.end_time else "В работе",
            duration,
            w.note or "",
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i + 1, column=col, value=val)
            cell.border = thin_border

    col_widths = [6, 20, 20, 12, 18, 18, 14, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"works_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
